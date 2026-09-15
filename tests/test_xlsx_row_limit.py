import datetime
import io
import unittest
from unittest import mock

import openpyxl
import xlsxwriter

from pyhindsight import analysis
from pyhindsight.analysis import AnalysisSession
from pyhindsight.browsers.webbrowser import WebBrowser

UTC = datetime.timezone.utc

# Two header rows plus four data rows per sheet, so ten records need three sheets.
MAX_ROWS = 6
RECORDS = 10


def _cookie(n):
    item = WebBrowser.CookieItem(
        profile='p', host_key=f'.site{n}.test', path='/', name=f'cookie-{n}', value=str(n),
        creation_utc=datetime.datetime(2024, 6, 1, 0, 0, n, tzinfo=UTC),
        last_access_utc=None, secure=False, http_only=False)
    item.interpretation = None
    item.source_item = 'Cookies'
    return item


def _local_storage(n):
    return WebBrowser.LocalStorageItem(
        profile='p', origin='https://example.test', key=f'key-{n}', value=str(n), seq=n,
        state='Live', source_path='Local Storage/leveldb', last_modified=None)


class TestXlsxRowLimit(unittest.TestCase):
    """Rows past a worksheet's last row continue on another sheet instead of vanishing.

    xlsxwriter returns -1 for a write past row 1,048,576 rather than raising, so the
    Timeline and Storage writers used to lose everything past the limit without a
    trace (#359). The limit is lowered here so the rollover can be seen with ten rows.
    """

    @classmethod
    def setUpClass(cls):
        session = AnalysisSession.__new__(AnalysisSession)
        session.parsed_artifacts = [_cookie(n) for n in range(RECORDS)]
        session.parsed_storage = [_local_storage(n) for n in range(RECORDS)]
        session.parsed_extension_data = []
        session.parsed_sync_data = []
        session.preferences = []
        session.installed_extensions = None
        session.plugin_results = {}
        session.timezone = UTC
        session.artifact_filter = None

        buffer = io.BytesIO()
        with mock.patch.object(analysis, 'XLSX_MAX_ROWS', MAX_ROWS), \
                mock.patch.object(analysis.log, 'warning') as cls.warnings:
            session.generate_excel(buffer)
        buffer.seek(0)
        cls.workbook = openpyxl.load_workbook(buffer)

    def _sheets(self, base):
        return [self.workbook[name] for name in self.workbook.sheetnames
                if name == base or name.startswith(f'{base} (')]

    def test_continuation_sheets_follow_the_sheet_they_continue(self):
        names = self.workbook.sheetnames
        self.assertEqual(['Timeline', 'Timeline (2)', 'Timeline (3)'], names[0:3])
        storage = names.index('Storage')
        self.assertEqual(['Storage', 'Storage (2)', 'Storage (3)'], names[storage:storage + 3])

    def test_every_record_is_written_exactly_once_in_order(self):
        for base, column, expected in (('Timeline', 4, [f'cookie-{n}' for n in range(RECORDS)]),
                                       ('Storage', 3, [f'key-{n}' for n in range(RECORDS)])):
            with self.subTest(sheet=base):
                values = [row[column - 1]
                          for sheet in self._sheets(base)
                          for row in sheet.iter_rows(min_row=3, values_only=True)]
                self.assertEqual(expected, values)

    def test_no_sheet_holds_more_rows_than_the_limit(self):
        for base in ('Timeline', 'Storage'):
            for sheet in self._sheets(base):
                with self.subTest(sheet=sheet.title):
                    self.assertLessEqual(sheet.max_row, MAX_ROWS)

    def test_continuations_repeat_the_header_widths_and_frozen_panes(self):
        for base in ('Timeline', 'Storage'):
            first, *continuations = self._sheets(base)
            header = [cell.value for cell in first[2]]
            for sheet in continuations:
                with self.subTest(sheet=sheet.title):
                    self.assertEqual(header, [cell.value for cell in sheet[2]])
                    self.assertEqual(first['A1'].value, sheet['A1'].value)
                    self.assertEqual(first.freeze_panes, sheet.freeze_panes)
                    self.assertEqual(first.column_dimensions['C'].width, sheet.column_dimensions['C'].width)

    def test_each_sheet_filters_its_own_rows(self):
        for base, last_column in (('Timeline', 'AG'), ('Storage', 'P')):
            refs = [sheet.auto_filter.ref for sheet in self._sheets(base)]
            with self.subTest(sheet=base):
                # Full sheets filter to their last row; the last one keeps the writer's
                # one-past-the-end row, as the single-sheet autofilter always did.
                self.assertEqual([f'A2:{last_column}6', f'A2:{last_column}6', f'A2:{last_column}5'], refs)

    def test_the_rollover_is_logged(self):
        messages = [call.args[0] for call in self.warnings.call_args_list]
        self.assertEqual(4, sum('continuing on' in message for message in messages), messages)


class TestXlsxRowLimitBoundary(unittest.TestCase):
    """At the real limit, the last row xlsxwriter accepts stays put and the next moves on."""

    def test_the_first_row_past_the_limit_starts_the_continuation(self):
        buffer = io.BytesIO()
        workbook = xlsxwriter.Workbook(buffer, {'in_memory': True})
        used = {'storage'}

        def continuation_name(base):
            name = f'{base} ({len(used) + 1})'
            used.add(name.lower())
            return name

        sheet = analysis.RolloverWorksheet(workbook, 'Storage', continuation_name)
        sheet.write(1, 0, 'Key')
        last_row = analysis.XLSX_MAX_ROWS - 1  # zero-based index of Excel row 1,048,576

        self.assertEqual(0, sheet.write(last_row, 0, 'last on the first sheet'))
        self.assertEqual(0, sheet.write(last_row + 1, 0, 'first on the continuation'))
        workbook.close()

        self.assertEqual(['Storage', 'Storage (2)'], [s.name for s in sheet.sheets])
        loaded = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()), read_only=True)
        continuation = list(loaded['Storage (2)'].iter_rows(values_only=True))
        self.assertEqual(('Key',), continuation[1])
        self.assertEqual(('first on the continuation',), continuation[2])


if __name__ == '__main__':
    unittest.main()
