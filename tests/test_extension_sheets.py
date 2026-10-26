import datetime
import io
import os
import types
import unittest

import openpyxl

from pyhindsight.analysis import AnalysisSession
from pyhindsight.browsers.firefox import Firefox
from pyhindsight.browsers.webbrowser import WebBrowser

UTC = datetime.timezone.utc
FIXTURE_DIR = os.path.join('tests', 'fixtures', 'firefox')


def _chrome_analysis():
    extension = WebBrowser.BrowserExtension(
        profile='chrome-profile', ext_id='aapocclcgogkmnckokdopfmhonfmgoek', name='Slides',
        description='Create and edit presentations', version='0.10', permissions=None,
        manifest='{}')
    return types.SimpleNamespace(installed_extensions={'data': [extension]},
                                 display_version='Chrome')


def _firefox_analysis():
    ff = Firefox(FIXTURE_DIR, no_copy=True, temp_dir=None, timezone=UTC)
    ff.get_extensions(FIXTURE_DIR, 'extensions.json')
    return types.SimpleNamespace(installed_extensions=ff.installed_extensions,
                                 display_version='Firefox')


class TestExtensionSheets(unittest.TestCase):
    """Both browsers' extensions reach the report once, on the Extensions worksheet.

    Firefox's used to carry a 'presentation', so the generic promotion also gave them
    a flat "Installed Extensions" sheet. In a run with a Chrome profile too, Chrome's
    extensions picked up that presentation and showed on the flat sheet as well.
    """

    @classmethod
    def setUpClass(cls):
        session = AnalysisSession()
        session.timezone = UTC
        session.installed_extensions = None
        session.absorb_browser_analysis('chrome-profile', _chrome_analysis())
        session.absorb_browser_analysis('firefox-profile', _firefox_analysis())
        cls.session = session

        buffer = io.BytesIO()
        session.generate_excel(buffer)
        buffer.seek(0)
        cls.workbook = openpyxl.load_workbook(buffer)

    def test_both_browsers_extensions_are_promoted_once(self):
        self.assertEqual(['data'], list(self.session.installed_extensions))
        ids = [ext.ext_id for ext in self.session.installed_extensions['data']]
        self.assertEqual(['aapocclcgogkmnckokdopfmhonfmgoek', 'uBlock0@raymondhill.net'], ids)

    def test_there_is_no_flat_installed_extensions_sheet(self):
        self.assertIn('Extensions', self.workbook.sheetnames)
        self.assertNotIn('Installed Extensions', self.workbook.sheetnames)

    def test_the_extension_header_rows_carry_version_and_firefox_details(self):
        headers = [row[0] for row in self.workbook['Extensions'].iter_rows(min_row=3, values_only=True)
                   if row[0] and '(' in row[0]]
        chrome, firefox = headers
        self.assertTrue(chrome.startswith('Slides (aapocclcgogkmnckokdopfmhonfmgoek) | Version: 0.10'))
        self.assertTrue(firefox.startswith('uBlock Origin (uBlock0@raymondhill.net) | Version: 1.55.0'))
        for part in ('State: Enabled', 'Found in: extensions.json', 'Install source: app-profile',
                     'Signing: Signed', 'Installed: 2024-01-15', 'Updated: 2024-01-22'):
            self.assertIn(part, firefox)


if __name__ == '__main__':
    unittest.main()
