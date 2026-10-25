import base64
import datetime
import email.utils
import hashlib
import io
import json
import os
import sqlite3
import tempfile
import unittest
import zlib
import zoneinfo

import openpyxl

from pyhindsight.analysis import AnalysisSession
from pyhindsight.browsers.chrome import Chrome
from pyhindsight.browsers.webbrowser import WebBrowser
from pyhindsight.lib import etag

UTC = datetime.timezone.utc
IST = zoneinfo.ZoneInfo('Asia/Kolkata')
# A fetch time after every sample below, so none is rejected as being from the future.
FETCHED = datetime.datetime(2026, 10, 24, tzinfo=UTC)


def utc(*args):
    return datetime.datetime(*args, tzinfo=UTC)


class TestRealWorldETags(unittest.TestCase):
    """Each format against values real servers sent, with the headers sent beside them.

    Captured in October 2026 with `curl -I`. Every decoded time is checked against the
    Last-Modified header from the same response and every decoded size against its
    Content-Length, so a sample here is evidence for the format, not an assumption
    about it.
    """

    # (server it came from, ETag, Last-Modified, Content-Length,
    #  expected format, expected time to the second, expected size)
    SAMPLES = [
        # The worked example in https://thelocalh0st.com/posts/etag-download-forensics/
        ('nginx (article)', '"6ac46e19-13e7f8"', 'Tue, 06 Oct 2026 03:42:17 GMT', 1304568,
         'nginx', utc(2026, 10, 6, 3, 42, 17), 1304568),
        ('www.nasa.gov', '"6acca3e6-c6"', 'Mon, 12 Oct 2026 09:09:58 GMT', 198,
         'nginx', utc(2026, 10, 12, 9, 9, 58), 198),
        ('www.python.org', '"6703bf29-1adb060"', 'Mon, 07 Oct 2024 10:59:53 GMT', 28160096,
         'nginx', utc(2024, 10, 7, 10, 59, 53), 28160096),
        ('www.apache.org', '"3c2e-6470630308aec"', 'Sun, 28 Dec 2025 17:09:24 GMT', 15406,
         'Apache httpd', utc(2025, 12, 28, 17, 9, 24), 15406),
        ('downloads.apache.org', '"86413-5caef2edf0372"', 'Wed, 01 Sep 2021 13:38:35 GMT', 549907,
         'Apache httpd', utc(2021, 9, 1, 13, 38, 35), 549907),
        ('the.earth.li (PuTTY)', '"1a0898-658fb8c515bde"', 'Fri, 14 Aug 2026 06:01:14 GMT', 1706136,
         'Apache httpd', utc(2026, 8, 14, 6, 1, 14), 1706136),
        ('live.sysinternals.com', '"a6c8bf9fa5cdd1:0"', 'Thu, 15 Oct 2026 23:14:48 GMT', 4776760,
         'IIS', utc(2026, 10, 15, 23, 14, 48), None),
        ('live.sysinternals.com', '"07e79518efedc1:0"', 'Wed, 17 Jun 2026 19:20:12 GMT', 1809728,
         'IIS', utc(2026, 6, 17, 19, 20, 12), None),
        ('live.sysinternals.com', '"09e80f3a06cd91:0"', 'Tue, 11 Apr 2023 18:10:52 GMT', 944520,
         'IIS', utc(2023, 4, 11, 18, 10, 52), None),
        ('live.sysinternals.com', '"0e495bbd8b4cc1:0"', 'Wed, 07 Dec 2011 12:07:04 GMT', 7903,
         'IIS', utc(2011, 12, 7, 12, 7, 4), None),
        ('download.sysinternals.com', '"0x8DF2B14CD9F2C17"', 'Thu, 15 Oct 2026 23:34:05 GMT', 5144350,
         'Azure Blob Storage', utc(2026, 10, 15, 23, 34, 5), None),
        ('GitHub release (Git for Windows)', '"0x8DCE7807154A6C6"', 'Tue, 08 Oct 2024 10:03:22 GMT',
         68778832, 'Azure Blob Storage', utc(2024, 10, 8, 10, 3, 22), None),
        ('learn.microsoft.com', 'W/"4316-1a141555da0"', 'Thu, 15 Oct 2026 20:51:16 GMT', 17174,
         'Express (Node.js)', utc(2026, 10, 15, 20, 51, 16), 17174),
        ('support.microsoft.com', '"1dd626648e2d1ee"', 'Thu, 22 Oct 2026 20:45:34 GMT', 15086,
         'ASP.NET Core', utc(2026, 10, 22, 20, 45, 34), None),
        ('www.irs.gov', '"1792680962"', 'Thu, 22 Oct 2026 14:56:02 GMT', None,
         'Unix timestamp', utc(2026, 10, 22, 14, 56, 2), None),
        ('www.sqlite.org (althttpd)', '"m67f42001s3153aa"', 'Mon, 07 Apr 2025 18:57:05 GMT', 3232682,
         'althttpd (SQLite)', utc(2025, 4, 7, 18, 57, 5), 3232682),
        ('www.7-zip.org', '"0x8DCBAC652C1AE4F"', 'Mon, 12 Aug 2024 12:00:13 GMT', 1624144,
         'Azure Blob Storage', utc(2024, 8, 12, 12, 0, 13), None),
        ('api.nuget.org', '0x8DB1FA95222F60F', 'Wed, 08 Mar 2023 07:47:09 GMT', 2441966,
         'Azure Blob Storage', utc(2023, 3, 8, 7, 47, 9), None),
    ]

    def test_each_sample_decodes_and_agrees_with_its_own_headers(self):
        for source, value, last_modified, size, fmt, modified, decoded_size in self.SAMPLES:
            with self.subTest(source=source, etag=value):
                result = etag.analyze(value, last_modified, size, reference_time=FETCHED)
                self.assertEqual(fmt, result.format)
                self.assertEqual(modified, result.modified.replace(microsecond=0))
                self.assertEqual(decoded_size, result.size)
                self.assertIn('time matches Last-Modified', result.notes)
                if decoded_size is not None:
                    self.assertIn('size matches total bytes', result.notes)

    def test_s3_multipart(self):
        # awscli.amazonaws.com/AWSCLIV2.msi: 50,204,672 bytes in 6 parts. The AWS CLI
        # uploads in 8 MiB parts, and 5 * 8 MiB < 50,204,672 <= 6 * 8 MiB.
        result = etag.analyze('"8584da86526b24ab6b47c4f45426f202-6"',
                              'Thu, 22 Oct 2026 19:05:14 GMT', 50204672)
        self.assertEqual('Amazon S3 multipart', result.format)
        self.assertIsNone(result.modified)
        self.assertIn('multipart upload in 6 parts; not an MD5 of the file', result.notes)
        self.assertIn('size consistent with 8 MiB parts', result.notes)

    def test_s3_multipart_with_more_than_one_possible_part_size(self):
        # rustup-init: 8 MiB parts, but 10 MiB parts would also make 3.
        result = etag.analyze('"9301526a1b7e8fe8bcce054d2e0a861d-3"', size=21139856)
        self.assertIn('size consistent with 8 or 10 MiB parts', result.notes)

    def test_md5(self):
        result = etag.analyze('"45b4cc5ffeb4ba28564a5872331d3d40"')  # nodejs.org
        self.assertEqual('MD5 hash', result.format)
        self.assertIsNone(result.modified)
        self.assertIsNone(result.size)

    def test_qiniu(self):
        # Live favicons; each hashes to its ETag.
        for source, value, last_modified, size in (
                ('www.qiniu.com', '"FsJyZYfBt-5Ve0VqMC4lKGTfRrGE"', 'Tue, 16 Aug 2022 07:04:11 GMT', 4286),
                ('dn-qiniu-avatar.qbox.me', '"FpFnkhvBzQRlgso9L3SxyxrZvppT"',
                 'Mon, 30 Mar 2015 14:34:04 GMT', 32988),
                ('www.wangsu.com', '"FlPYJRi-se1kdokEJJORKD_4RFdU"', 'Sun, 09 Aug 2020 03:23:35 GMT', 4286)):
            with self.subTest(source=source):
                result = etag.analyze(value, last_modified, size)
                self.assertEqual('Qiniu Kodo', result.format)
                self.assertIsNone(result.modified)
                self.assertEqual(['SHA-1 of the content; the file is 4 MiB or less'], result.notes)

    def test_a_hash_shape_alone_is_not_named_as_a_hash(self):
        # docs.aws.amazon.com is a SHA-1, learn.microsoft.com is not; GitHub avatars vary.
        for value in ('"9ATiBeUUEYYreZx9C3jOaqQBVHo="', '"SSctglCkYMnDerQO6mEzUN5Gsrk="',
                      '"2f6e95142316814ab3cebc6f656728bdc0315f23000c417222c763b8c49ea94d"'):
            with self.subTest(etag=value):
                self.assertEqual('Unrecognized', etag.analyze(value, reference_time=FETCHED).format)

    def test_netlify_is_not_read_as_an_s3_part_count(self):
        # Netlify; -ssl is not a part count.
        for value in ('"e54964d35ed877151c08a3621b3d415a-ssl"',
                      'W/"e0788bff0f1ad1cdca751fab8e3b85d4-ssl-df"'):
            with self.subTest(etag=value):
                self.assertEqual('Unrecognized', etag.analyze(value).format)

    def test_a_long_0x_value_is_not_read_as_azure_ticks(self):
        # download.visualstudio.microsoft.com; the file's SHA-256 is e38048c4...e58b.
        result = etag.analyze('"0x8BC5ECAEE96421C22120ACD02B30807A560F31C892E13FC63EADECC3AC6B9C1A"',
                              'Wed, 07 Oct 2026 05:02:07 GMT', 4475728, reference_time=FETCHED)
        self.assertEqual('Microsoft Distribute (0x + 256 bits)', result.format)
        self.assertIsNone(result.modified)
        self.assertIsNone(result.size)
        self.assertEqual(['not a time and not the file\'s SHA-256; compare the whole value with a '
                          'later download of the same file'], result.notes)


class TestAkamaiNetStorage(unittest.TestCase):
    """JavaSetup8u351.exe is from magnet.ctf_2023; the rest were captured live."""

    JAVA_SETUP = '"82bc7b7e2716e6a631952daa1be4037e:1668626955.19894"'
    JAVA_SETUP_LAST_MODIFIED = 'Wed, 16 Nov 2022 19:00:10 GMT'

    def test_the_download_record(self):
        result = etag.analyze(self.JAVA_SETUP, self.JAVA_SETUP_LAST_MODIFIED, 2261648,
                              reference_time=utc(2023, 1, 4, 0, 26, 39))
        self.assertEqual('Akamai NetStorage-style', result.format)
        self.assertFalse(result.weak)
        # The suffix is 29 minutes after Last-Modified.
        self.assertIsNone(result.modified)
        self.assertIsNone(result.size)
        self.assertEqual(['usually the MD5 of the content; compare with the MD5 of a recovered copy',
                          'suffix 1668626955.19894 is not read as the modification time; as a '
                          'Unix time it is 0:29:05 after Last-Modified'], result.notes)

    def test_the_prefix_is_the_md5_of_the_body(self):
        # download.oracle.com: the 64-byte body hashes to the prefix.
        body = b'6deafad80454d1f968105b1f582d89b8c2e88a5dfc98116a8226e0ec2fd5df8e'
        value = '"b670c22a9f59b6f8839abcb860d4bbe6:1792846992.069644"'
        result = etag.analyze(value, 'Sat, 24 Oct 2026 12:41:06 GMT', len(body), body=body)
        self.assertEqual('Akamai NetStorage-style', result.format)
        self.assertIn('matches the MD5 of the cached body', result.notes)
        self.assertIn('suffix 1792846992.069644 is not read as the modification time; as a '
                      'Unix time it is 0:22:06 after Last-Modified', result.notes)
        self.assertIn('does not match the MD5 of the cached body',
                      etag.analyze(value, body=body[:-1]).notes)

    def test_a_suffix_on_the_second_of_last_modified_is_still_not_the_time(self):
        # www.salesforce.com
        result = etag.analyze('"ccda8df05e9a37b3131afd4d451b44ec:1765795051.090913"',
                              'Mon, 15 Dec 2025 10:37:31 GMT', 5430)
        self.assertIsNone(result.modified)
        self.assertIn('suffix 1765795051.090913 is not read as the modification time; as a '
                      'Unix time it is the second of Last-Modified', result.notes)

    def test_a_vendor_published_md5(self):
        # VirtualBox 7.0.18: the prefix is the MD5 in Oracle's MD5SUMS.
        result = etag.analyze('"6a046a57ca3dd222d8bf1410b8172f81:1714736344.35986"',
                              'Thu, 02 May 2024 11:59:19 GMT', 109700144)
        self.assertEqual('Akamai NetStorage-style', result.format)
        self.assertIsNone(result.modified)
        self.assertIn('suffix 1714736344.35986 is not read as the modification time; as a '
                      'Unix time it is 23:39:45 after Last-Modified', result.notes)

    def test_a_suffix_days_after_last_modified(self):
        # yum.oracle.com/RPM-GPG-KEY-oracle-ol9
        result = etag.analyze('"57d39f1538ab7dcd46f2c8662f55a96f:1657036179.196558"',
                              'Wed, 29 Jun 2022 20:50:51 GMT', 6350)
        self.assertIsNone(result.modified)
        self.assertIn('suffix 1657036179.196558 is not read as the modification time; as a '
                      'Unix time it is 5 days, 18:58:48 after Last-Modified', result.notes)

    def test_without_last_modified(self):
        result = etag.analyze('"402e7a087747cb56c718bde84651f96a:1367999472"')
        self.assertEqual('Akamai NetStorage-style', result.format)
        self.assertIn('suffix 1367999472 is not read as the modification time', result.notes)
        self.assertFalse(any(note.startswith('not confirmed') for note in result.notes))

    def test_weak_and_unquoted_values(self):
        weak = etag.analyze('W/' + self.JAVA_SETUP)
        self.assertEqual('Akamai NetStorage-style', weak.format)
        self.assertTrue(weak.weak)
        self.assertEqual('Akamai NetStorage-style', etag.analyze(self.JAVA_SETUP.strip('"')).format)

    def test_near_misses_are_not_claimed(self):
        for value in ('"82bc7b7e2716e6a631952daa1be4037e:1668626955.1.2"',
                      '"82bc7b7e2716e6a631952daa1be4037:1668626955"',
                      '"82bc7b7e2716e6a631952daa1be4037e:"',
                      '"82bc7b7e2716e6a631952daa1be4037e:-1668626955"',
                      '"82bc7b7e2716e6a631952daa1be4037e:0x63753d0b"'):
            with self.subTest(etag=value):
                self.assertEqual('Unrecognized', etag.analyze(value).format)


class TestCorpusFindings(unittest.TestCase):
    """Cases found by running the decoder over the bf4sa_2025_bob-1 corpus root."""

    def test_a_precompressed_file_matches_the_compressed_length(self):
        # An Express origin serving .js.gz files: the ETag's size is the gzip file's
        # (501, the Content-Length), while the cache holds the decoded 963 bytes.
        mtime = datetime.datetime.fromtimestamp(0x1967e7e9c04 / 1000, UTC)
        result = etag.analyze('W/"1f5-1967e7e9c04"', email.utils.format_datetime(mtime, usegmt=True),
                              size=963, size_label='the cached body size',
                              content_encoding='gzip', encoded_size=501)
        self.assertEqual('Express (Node.js)', result.format)
        self.assertIn('size matches the gzip-compressed Content-Length', result.notes)

    def test_a_size_matching_neither_length_names_both(self):
        # The same origin behind a CDN that recompressed it as brotli.
        mtime = datetime.datetime.fromtimestamp(0x1967e7e9c04 / 1000, UTC)
        result = etag.analyze('W/"219-1967e7e9c04"', email.utils.format_datetime(mtime, usegmt=True),
                              size=799, size_label='the cached body size',
                              content_encoding='br', encoded_size=486)
        self.assertEqual('Express (Node.js)', result.format)
        self.assertIn('size differs from the cached body size (799) and the br-compressed '
                      'Content-Length (486)', result.notes)

    def test_a_hash_mismatch_on_a_compressed_response_says_why_it_might_be(self):
        value = hashlib.md5(b'hello').hexdigest()
        result = etag.analyze(value, body=b'jello', content_encoding='gzip', encoded_size=25)
        self.assertIn('does not match the MD5 of the cached body (decoded from gzip; the ETag '
                      'may describe the compressed bytes)', result.notes)

    def test_a_precompressed_file_matches_the_md5_of_the_bytes_as_sent(self):
        # The ETag of a precompressed file is the MD5 of the compressed bytes.
        decoded, sent = b'hello', b'\x0b\x02\x80hello\x03'
        value = hashlib.md5(sent).hexdigest()
        result = etag.analyze(value, body=decoded, content_encoding='br',
                              encoded_md5=hashlib.md5(sent).hexdigest())
        self.assertIn('matches the MD5 of the cached body as sent (br-compressed)', result.notes)
        self.assertIn('does not match the MD5 of the cached body (decoded from br; the ETag '
                      'may describe the compressed bytes)',
                      etag.analyze(value, body=decoded, content_encoding='br').notes)

    def test_upper_case_filetime(self):
        # www.montereybayaquarium.org (Optimizely CMS); every one matched Last-Modified.
        result = etag.analyze('"1DA771657FC7680"', 'Fri, 15 Mar 2024 20:21:21 GMT')
        self.assertEqual('.NET (FILETIME)', result.format)
        self.assertEqual(utc(2024, 3, 15, 20, 21, 21), result.modified.replace(microsecond=0))
        unconfirmed = etag.analyze('"1DA771657FC7680"')
        self.assertEqual('.NET (FILETIME) (unconfirmed)', unconfirmed.format)
        self.assertIsNone(unconfirmed.modified)


class TestDecoding(unittest.TestCase):

    def test_no_etag(self):
        for value in (None, '', '   ', b'"6ac46e19-13e7f8"', 12345):
            with self.subTest(value=value):
                self.assertIsNone(etag.analyze(value))

    def test_quotes_are_optional(self):
        # Chrome stores the header as the server sent it, and some servers omit the
        # quotes the RFC requires (upload.wikimedia.org among them).
        quoted = etag.analyze('"6ac46e19-13e7f8"')
        bare = etag.analyze('6ac46e19-13e7f8')
        self.assertEqual((quoted.format, quoted.modified, quoted.size),
                         (bare.format, bare.modified, bare.size))

    def test_weak_etags_still_decode_and_say_they_are_weak(self):
        # nginx turns its ETag weak when it gzips the response, and keeps the value.
        result = etag.analyze('W/"6ac46e19-13e7f8"')
        self.assertEqual('nginx', result.format)
        self.assertTrue(result.weak)
        self.assertEqual(1304568, result.size)
        self.assertTrue(result.notes[0].startswith('weak ETag'))

    def test_an_unknown_shape_is_unrecognized_not_guessed(self):
        # dl.google.com: hyphenated hex, but no reading of it gives a plausible time.
        result = etag.analyze('"6e71e56-e4d464d1"', 'Mon, 12 Oct 2026 20:47:03 GMT', 13283768)
        self.assertEqual('Unrecognized', result.format)
        self.assertIsNone(result.modified)
        self.assertIsNone(result.size)

    def test_a_decode_every_check_contradicts_is_discarded(self):
        # Shaped like nginx, but neither the time nor the size agrees with the response
        # it came with, so the shape matched by accident.
        result = etag.analyze('"6ac46e19-13e7f8"', 'Mon, 01 Jan 2024 00:00:00 GMT', 999)
        self.assertEqual('Unrecognized', result.format)

    def test_a_single_failed_check_is_reported_rather_than_hidden(self):
        # The time agrees, so the format is credible, and the size disagreeing is a
        # finding in its own right (a different copy, or a transformed response).
        result = etag.analyze('"6ac46e19-13e7f8"', 'Tue, 06 Oct 2026 03:42:17 GMT', 1000)
        self.assertEqual('nginx', result.format)
        self.assertIn('time matches Last-Modified', result.notes)
        self.assertIn('size differs from total bytes (1,000)', result.notes)

    def test_a_time_mismatch_says_by_how_much(self):
        result = etag.analyze('"6ac46e19-13e7f8"', 'Tue, 06 Oct 2026 05:42:17 GMT', 1304568)
        self.assertEqual('nginx', result.format)
        self.assertIn('time is 2:00:00 before Last-Modified', result.notes)

    def test_shapes_that_need_last_modified_are_only_a_reading_without_it(self):
        # Common enough to match by chance, so only a reading without Last-Modified.
        for value, fmt, reading in (
                ('"1dd626648e2d1ee"', 'ASP.NET Core (unconfirmed)', '2026-10-22 20:45:34 UTC'),
                ('"1792680962"', 'Unix timestamp (unconfirmed)', '2026-10-22 14:56:02 UTC'),
                # i.ytimg.com
                ('"1749462010"', 'Unix timestamp (unconfirmed)', '2025-06-09 09:40:10 UTC')):
            with self.subTest(etag=value):
                result = etag.analyze(value, size=15086)
                self.assertEqual(fmt, result.format)
                self.assertIsNone(result.modified)
                self.assertEqual([f'reads as {reading}; not confirmed: no Last-Modified to '
                                  f'check against'], result.notes)

    def test_shapes_that_last_modified_contradicts_are_dropped(self):
        for value in ('"1dd626648e2d1ee"', '"1792680962"'):
            with self.subTest(etag=value):
                self.assertEqual('Unrecognized',
                                 etag.analyze(value, 'Mon, 01 Jan 2024 00:00:00 GMT', 15086).format)

    def test_values_that_do_not_read_as_a_plausible_time_stay_unrecognized(self):
        for value in ('"26177639"', '"4710585"', '8390199704425924230', '"9999999999"'):
            with self.subTest(etag=value):
                self.assertEqual('Unrecognized', etag.analyze(value, reference_time=FETCHED).format)

    def test_the_gap_before_the_download_is_reported(self):
        result = etag.analyze('"6ac46e19-13e7f8"', reference_time=utc(2026, 10, 7, 3, 42, 17))
        self.assertIn('modified 1 day, 0:00:00 before the download', result.notes)

    def test_a_time_after_the_download_is_flagged(self):
        result = etag.analyze('"6ac46e19-13e7f8"', reference_time=utc(2026, 10, 6, 3, 0, 0))
        self.assertIn('modified 0:42:17 after the download (server and browser clocks disagree?)',
                      result.notes)

    def test_a_time_far_after_the_download_is_not_a_timestamp(self):
        result = etag.analyze('"6ac46e19-13e7f8"', reference_time=utc(2020, 1, 1))
        self.assertEqual('Unrecognized', result.format)

    def test_a_decode_with_nothing_to_check_it_against_says_so(self):
        unchecked = etag.analyze('"6ac46e19-13e7f8"', reference_time=FETCHED)
        self.assertEqual('nginx', unchecked.format)
        self.assertIn('not confirmed: no Last-Modified or size recorded to check against',
                      unchecked.notes)
        for checked in (etag.analyze('"6ac46e19-13e7f8"', 'Tue, 06 Oct 2026 03:42:17 GMT'),
                        etag.analyze('"6ac46e19-13e7f8"', size=1304568)):
            with self.subTest(notes=checked.notes):
                self.assertFalse(any(note.startswith('not confirmed') for note in checked.notes))

    def test_iis_without_last_modified_lists_the_possibilities(self):
        # Without Last-Modified the unpadded bytes split several ways. The ETag is still
        # identifiably IIS; the time is not, and saying so beats picking one.
        result = etag.analyze('"a6c8bf9fa5cdd1:0"', reference_time=FETCHED, timezone=IST)
        self.assertEqual('IIS', result.format)
        self.assertIsNone(result.modified)
        self.assertEqual(1, len(result.notes))
        self.assertIn('possible timestamps from 2026-07-05 23:12:18 IST', result.notes[0])
        self.assertIn('Last-Modified is needed to choose', result.notes[0])

    def test_iis_change_number(self):
        result = etag.analyze('"a6c8bf9fa5cdd1:3"', 'Thu, 15 Oct 2026 23:14:48 GMT')
        self.assertEqual('IIS', result.format)
        self.assertIn('configuration change number 3', result.notes)

    def test_apache_with_inode_and_a_compression_suffix(self):
        result = etag.analyze('"1a2b3c-3c2e-6470630308aec-gzip"', 'Sun, 28 Dec 2025 17:09:24 GMT')
        self.assertEqual('Apache httpd', result.format)
        self.assertEqual(15406, result.size)
        self.assertEqual(utc(2025, 12, 28, 17, 9, 24), result.modified.replace(microsecond=0))
        self.assertIn(f'inode {0x1a2b3c}', result.notes)
        self.assertIn('compressed variant (-gzip); size is of the uncompressed file', result.notes)

    def test_tomcat(self):
        modified = utc(2025, 12, 28, 17, 9, 24)
        value = f'W/"15406-{int(modified.timestamp()) * 1000}"'
        result = etag.analyze(value, 'Sun, 28 Dec 2025 17:09:24 GMT', 15406)
        self.assertEqual('Tomcat', result.format)
        self.assertEqual(modified, result.modified)
        self.assertEqual(15406, result.size)

    def test_express_body_hash_is_checked_against_a_cached_body(self):
        body = b'hello'
        digest = base64.b64encode(hashlib.sha1(body).digest()).decode()[:27]
        value = f'W/"{len(body):x}-{digest}"'
        matching = etag.analyze(value, size=len(body), body=body)
        self.assertEqual('Express (Node.js)', matching.format)
        self.assertEqual(5, matching.size)
        self.assertIn('matches the cached body', matching.notes)
        self.assertIn('does not match the cached body',
                      etag.analyze(value, body=b'jello').notes)

    def test_md5_is_checked_against_a_cached_body(self):
        value = hashlib.md5(b'hello').hexdigest()
        self.assertIn('matches the MD5 of the cached body',
                      etag.analyze(value, body=b'hello').notes)
        self.assertIn('does not match the MD5 of the cached body',
                      etag.analyze(value, body=b'jello').notes)

    def test_werkzeug(self):
        # Werkzeug 2.x send_file; 3.x drops the checksum.
        path = b'/srv/app/static/report.pdf'
        check = zlib.adler32(path)
        for value, notes in ((f'"1718794392.4765725-1874910-{check}"',
                              [f'adler32 of the file\'s path on the server {check}']),
                             ('"1718794392.4765725-1874910"', [])):
            with self.subTest(etag=value):
                result = etag.analyze(value, 'Wed, 19 Jun 2024 10:53:12 GMT', 1874910)
                self.assertEqual('Werkzeug (Flask)', result.format)
                self.assertEqual(datetime.datetime(2024, 6, 19, 10, 53, 12, 476572, tzinfo=UTC),
                                 result.modified)
                self.assertEqual(1874910, result.size)
                self.assertEqual(notes + ['time matches Last-Modified', 'size matches total bytes'],
                                 result.notes)

    def test_a_reproducible_build_time_is_accepted_when_last_modified_agrees(self):
        # Reproducible builds stamp files 1980-01-01.
        value = f'W/"1bacd64-{315532801000:x}"'  # Express: size, mtime in milliseconds
        result = etag.analyze(value, 'Tue, 01 Jan 1980 00:00:01 GMT', 29019492)
        self.assertEqual('Express (Node.js)', result.format)
        self.assertEqual(utc(1980, 1, 1, 0, 0, 1), result.modified)
        self.assertIn('time matches Last-Modified', result.notes)
        self.assertEqual('Unrecognized', etag.analyze(value, size=29019492).format)
        self.assertEqual('Unrecognized',
                         etag.analyze(value, 'Tue, 01 Jan 1980 00:00:05 GMT', 29019492).format)

    def test_qiniu_is_checked_against_a_cached_body(self):
        def qetag(prefix, digest):
            return '"' + base64.urlsafe_b64encode(prefix + digest).decode() + '"'

        small = b'hello'
        value = qetag(b'\x16', hashlib.sha1(small).digest())
        self.assertIn('matches the cached body', etag.analyze(value, body=small).notes)
        self.assertIn('does not match the cached body', etag.analyze(value, body=b'jello').notes)

        # Over 4 MiB, the SHA-1 is of the blocks' SHA-1s.
        block = 4 * 1024 * 1024
        large = b'a' * block + b'b'
        value = qetag(b'\x96', hashlib.sha1(hashlib.sha1(large[:block]).digest()
                                            + hashlib.sha1(large[block:]).digest()).digest())
        result = etag.analyze(value, size=len(large), body=large)
        self.assertIn('SHA-1 of the 4 MiB blocks\' SHA-1s; the file is over 4 MiB', result.notes)
        self.assertIn('matches the cached body', result.notes)
        self.assertIn('the recorded size is on the other side of 4 MiB',
                      etag.analyze(value, size=len(small)).notes)

    def test_malformed_values_do_not_raise(self):
        for value in ('W/', '"', '""', 'W/""', '0x', '0xFFFFFFFFFFFFFFFF', 'f' * 16 + ':0',
                      'ffffffffffffffff-ffffffffffffffff', '-', ':', '9' * 13 + '-' + '9' * 13,
                      '\x00\xff', '"' * 50):
            with self.subTest(etag=value):
                etag.analyze(value, 'not a date', -1, reference_time=FETCHED, body=b'x')


def _download(**overrides):
    fields = dict(
        etag='"6ac46e19-13e7f8"', last_modified='Tue, 06 Oct 2026 03:42:17 GMT',
        total_bytes=1304568, start_time=utc(2026, 10, 7, 3, 42, 17))
    fields.update(overrides)
    item = WebBrowser.DownloadItem(
        'Default', 1, 'https://example.test/payload.exe', fields['total_bytes'],
        fields['total_bytes'], 1, start_time=fields['start_time'], etag=fields['etag'],
        last_modified=fields['last_modified'], opened=0)
    item.row_type = 'download'
    item.timestamp = item.start_time
    item.status_friendly = 'Complete'
    item.value = r'C:\Users\x\Downloads\payload.exe'
    item.interpretation = ''
    item.source_item = 'History'
    item.interrupt_reason_friendly = 'No Interrupt'
    item.danger_type_friendly = 'Not Dangerous'
    return item


def _session(items, timezone=UTC):
    session = AnalysisSession.__new__(AnalysisSession)
    session.parsed_artifacts = items
    session.parsed_storage = []
    session.parsed_extension_data = []
    session.parsed_sync_data = []
    session.preferences = []
    session.installed_extensions = None
    session.plugin_results = {}
    session.artifact_filter = None
    session.timezone = timezone
    return session


class TestSessionAnalysis(unittest.TestCase):

    def test_a_download_is_decoded_into_the_display_timezone(self):
        item = _download(start_time=utc(2026, 10, 7, 3, 42, 17).astimezone(IST))
        _session([item], IST).apply_etag_analysis()
        self.assertEqual('nginx', item.etag_format)
        self.assertEqual(1304568, item.etag_size)
        self.assertEqual(datetime.datetime(2026, 10, 6, 9, 12, 17, tzinfo=IST), item.etag_modified)
        self.assertEqual(datetime.timedelta(hours=5, minutes=30), item.etag_modified.utcoffset())
        self.assertEqual('time matches Last-Modified; size matches total bytes; '
                         'modified 1 day, 0:00:00 before the download', item.etag_notes)

    def test_a_cached_body_is_used_only_when_it_is_the_resource(self):
        body = b'hello'
        item = WebBrowser.CacheItem('Default', 'https://example.test/a', None,
                                    FETCHED, 'locations', 'key', None, body)
        item.row_type = 'cache'
        item.etag = hashlib.md5(body).hexdigest()

        _session([item]).apply_etag_analysis()
        # No body_sha256 means the bytes may still be compressed, so they are not
        # compared: a mismatch against a gzip stream would be a false finding.
        self.assertNotIn('MD5 of the cached body', item.etag_notes)

        item.body_sha256 = hashlib.sha256(body).hexdigest()
        _session([item]).apply_etag_analysis()
        self.assertIn('matches the MD5 of the cached body', item.etag_notes)

    def test_a_range_response_is_checked_against_the_full_size(self):
        # A 206 holds a slice of the file, so the slice's length and hash are not the
        # file's. Content-Range carries the full size after the slash.
        class Metadata:
            def __init__(self, headers):
                self.headers = headers

            def get_attribute(self, name):
                return [self.headers[name]] if name in self.headers else []

        part = b'x' * 100
        for content_range, size, note in (
                ('bytes 0-99/1304568', 1304568, 'size matches the Content-Range total'),
                ('bytes 0-99/*', None, 'not confirmed')):
            with self.subTest(content_range=content_range):
                item = WebBrowser.CacheItem(
                    'Default', 'https://example.test/a', None, FETCHED, 'locations', 'key',
                    Metadata({'content-range': content_range, 'content-length': '100'}), part)
                item.row_type = 'cache'
                item.etag = '"6ac46e19-13e7f8"'
                item.body_sha256 = hashlib.sha256(part).hexdigest()
                _session([item]).apply_etag_analysis()
                self.assertEqual('nginx', item.etag_format)
                self.assertEqual(1304568, item.etag_size)
                self.assertIn(note, item.etag_notes)
                self.assertNotIn('size differs', item.etag_notes)

    def test_a_compressed_cache_body_is_also_checked_as_sent(self):
        class Metadata:
            def get_attribute(self, name):
                return {'content-encoding': ['br'], 'content-length': ['12']}.get(name, [])

        decoded, sent = b'hello', b'\x0b\x02\x80hello\x03'
        item = WebBrowser.CacheItem('Default', 'https://example.test/a.js', None, FETCHED,
                                    'locations', 'key', Metadata(), decoded)
        item.row_type = 'cache'
        item.etag = f'"{hashlib.md5(sent).hexdigest()}"'
        item.body_sha256 = hashlib.sha256(decoded).hexdigest()
        item.encoded_md5 = hashlib.md5(sent).hexdigest()
        _session([item]).apply_etag_analysis()
        self.assertEqual('MD5 hash', item.etag_format)
        self.assertIn('matches the MD5 of the cached body as sent (br-compressed)', item.etag_notes)

    def test_items_without_an_etag_are_left_alone(self):
        item = _download(etag=None)
        _session([item]).apply_etag_analysis()
        self.assertIsNone(item.etag_format)

    def test_real_history_downloads(self):
        # upload.wikimedia.org serves from OpenStack Swift, whose ETag is the object's
        # MD5. Hard_disk.jpg was downloaded into the v30 profile in 2013 and the v60 one
        # in 2017: Last-Modified moved between the two, the ETag did not, and the file
        # it still serves hashes to that same value. That is the correlation this
        # decoding exists to support.
        decoded = {}
        for version in (30, 60):
            browser = Chrome(os.path.join('tests', 'fixtures', 'profiles', str(version)),
                             version=[version], no_copy=True, timezone=UTC)
            browser.get_downloads(browser.profile_path, 'History', browser.version, 'download')
            _session(browser.parsed_artifacts).apply_etag_analysis()
            hard_disk = next(item for item in browser.parsed_artifacts
                             if item.url.endswith('Hard_disk.jpg'))
            decoded[version] = hard_disk

        for item in decoded.values():
            self.assertEqual('ea6b7682ed897a3e9cf22274e431a455', item.etag)
            self.assertEqual('MD5 hash', item.etag_format)
        self.assertNotEqual(decoded[30].last_modified, decoded[60].last_modified)


class TestOutputs(unittest.TestCase):
    """The decoded fields reach all three output formats, under their own names."""

    def setUp(self):
        self.directory = tempfile.mkdtemp()
        self.session = _session([_download()], IST)
        self.session.apply_etag_analysis()

    def tearDown(self):
        for name in os.listdir(self.directory):
            os.remove(os.path.join(self.directory, name))
        os.rmdir(self.directory)

    def test_xlsx(self):
        buffer = io.BytesIO()
        self.session.generate_excel(buffer)
        buffer.seek(0)
        sheet = openpyxl.load_workbook(buffer)['Timeline']
        row = {header.value: value.value for header, value in zip(sheet[2], sheet[3])}

        self.assertEqual('"6ac46e19-13e7f8"', row['ETag'])
        # One readable cell; the parts are separate fields in SQLite and JSONL.
        self.assertEqual('nginx | modified 2026-10-06 09:12:17.000 | 1,304,568 bytes | '
                         'time matches Last-Modified; size matches total bytes; '
                         'modified 1 day, 0:00:00 before the download',
                         row['ETag Decoded (Asia/Kolkata)'])
        # The columns after the inserted one moved right and still hold their own data.
        self.assertIn('All HTTP Headers', row)
        self.assertIn('Body SHA256', row)
        self.assertEqual('AH', sheet.auto_filter.ref.split(':')[1].rstrip('0123456789'))

    def test_xlsx_leaves_the_cell_empty_without_an_etag(self):
        self.session.parsed_artifacts = [_download(etag=None)]
        self.session.apply_etag_analysis()
        buffer = io.BytesIO()
        self.session.generate_excel(buffer)
        buffer.seek(0)
        sheet = openpyxl.load_workbook(buffer)['Timeline']
        row = {header.value: value.value for header, value in zip(sheet[2], sheet[3])}
        self.assertIsNone(row['ETag Decoded (Asia/Kolkata)'])

    def test_xlsx_leaves_the_cell_empty_for_an_unrecognized_etag(self):
        # SQLite and JSONL still say 'Unrecognized'; the XLSX cell would only repeat
        # the ETag column beside it.
        self.session.parsed_artifacts = [_download(etag='"6e71e56-e4d464d1"')]
        self.session.apply_etag_analysis()
        self.assertEqual('Unrecognized', self.session.parsed_artifacts[0].etag_format)
        buffer = io.BytesIO()
        self.session.generate_excel(buffer)
        buffer.seek(0)
        sheet = openpyxl.load_workbook(buffer)['Timeline']
        row = {header.value: value.value for header, value in zip(sheet[2], sheet[3])}
        self.assertIsNone(row['ETag Decoded (Asia/Kolkata)'])

    def test_sqlite(self):
        path = os.path.join(self.directory, 'out.sqlite')
        self.session.generate_sqlite(path)
        con = sqlite3.connect(path)
        try:
            self.assertEqual(
                ('nginx', '2026-10-06 09:12:17.000', 1304568),
                con.execute('SELECT etag_format, etag_modified, etag_size FROM timeline').fetchone())
            # Appended after the existing columns, so their positions are unchanged.
            columns = [row[1] for row in con.execute('PRAGMA table_info(timeline)')]
            self.assertEqual(['body_sha256', 'etag_format', 'etag_modified', 'etag_size',
                              'etag_notes'], columns[-5:])
        finally:
            con.close()

    def test_jsonl(self):
        path = os.path.join(self.directory, 'out.jsonl')
        self.session.generate_jsonl(path)
        with open(path, encoding='utf-8') as f:
            record = json.loads(f.readline())
        self.assertEqual('nginx', record['etag_format'])
        self.assertEqual('2026-10-06T09:12:17+05:30', record['etag_modified'])
        self.assertEqual(1304568, record['etag_size'])
        self.assertIn('size matches total bytes', record['etag_notes'])


if __name__ == '__main__':
    unittest.main()
